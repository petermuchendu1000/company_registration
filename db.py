"""
Supabase (PostgreSQL) data-access layer — the single source of truth for the
company pipeline. Replaces the old Excel-backed store.

All rows are exchanged as dicts keyed by the SAME human header labels the rest
of the codebase already uses (e.g. "Company Number", "Company Name",
"Date of Creation"), so callers don't need to know about SQL columns.

Connection string is read from the environment (DATABASE_URL, falling back to
SUPABASE_DB_URL). Never hard-code credentials — put them in .env (git-ignored).

Public API:
    ensure_schema()                       -> create the companies table if absent
    all_rows()                            -> list[dict] keyed by header labels
    find_row(cn)                          -> dict | None
    add_company(cn, name, date="")        -> bool (True if inserted)
    update_row(cn, {header: value, ...})  -> None (upserts the company if new)
    save_results(results)                 -> int (upsert nested pipeline results)
    import_from_excel(path)               -> int (one-time migration helper)
    company_numbers()                     -> set[str]
"""
from __future__ import annotations

import os
import threading
import time

import psycopg2
import psycopg2.extras

try:
    from dotenv import load_dotenv
    load_dotenv()
except Exception:
    pass


# ── Column mapping ──────────────────────────────────────────────────────────
# (human header label, snake_case db column). Order defines display/export order.
COLUMNS: list[tuple[str, str]] = [
    ("No.",                     "row_no"),
    ("Company Number",          "company_number"),
    ("Company Name",            "company_name"),
    ("Status",                  "status"),
    ("Type",                    "type"),
    ("Date of Creation",        "date_of_creation"),
    ("SIC Codes",               "sic_codes"),
    ("Address",                 "address"),
    ("Directors",               "directors"),
    ("Director Nationalities",  "director_nationalities"),
    ("Director DOB",            "director_dob"),
    ("Director Gender",         "director_gender"),
    ("Buvei First Name",        "buvei_first_name"),
    ("Buvei Last Name",         "buvei_last_name"),
    ("DUNS Number",             "duns_number"),
    ("DUNS Status",             "duns_status"),
    ("DUNS Email Used",         "duns_email_used"),
    ("D&B Legal Name",          "dnb_legal_name"),
    ("D&B Address",             "dnb_address"),
    ("Certificate Downloaded",  "certificate_downloaded"),
    ("Certificate Path",        "certificate_path"),
    ("Domain",                  "domain"),
    ("Domain Status",           "domain_status"),
    ("Domain Cost",             "domain_cost"),
    ("Assigned Email",          "assigned_email"),
    ("Account Name",            "account_name"),
    ("Developer Email",         "developer_email"),
    ("Dev Email Forwarding",    "dev_email_forwarding"),
    ("Google TXT Status",       "google_txt_status"),
    ("Organization Phone",      "organization_phone"),
    ("Play Signup Missing",     "play_signup_missing"),
    ("Company Status",          "pipeline_status"),
    ("Company Notes",           "company_notes"),
    ("Archived",                "archived"),
]
HEADER_TO_COL = {h: c for h, c in COLUMNS}
COL_TO_HEADER = {c: h for h, c in COLUMNS}
TABLE = "companies"


# ── Connection ──────────────────────────────────────────────────────────────
def _dsn() -> str:
    dsn = os.environ.get("DATABASE_URL") or os.environ.get("SUPABASE_DB_URL")
    if not dsn:
        raise RuntimeError(
            "DATABASE_URL not set. Add it to .env "
            "(postgresql://...pooler.supabase.com:5432/postgres)."
        )
    return dsn


def _connect():
    return psycopg2.connect(_dsn(), connect_timeout=15,
                            cursor_factory=psycopg2.extras.RealDictCursor)


# ── Value normalisation ─────────────────────────────────────────────────────
def _cn8(v) -> str:
    s = str(v or "").strip()
    return s.zfill(8) if s.isdigit() else s


def _norm(value):
    """Coerce a Python value into what we store (text, or None)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "Yes" if value else ""
    if isinstance(value, (int, float)):
        return str(value)
    s = str(value).strip()
    return s if s != "" else None


# ── Schema ──────────────────────────────────────────────────────────────────
def ensure_schema() -> None:
    cols_sql = []
    for header, col in COLUMNS:
        if col == "company_number":
            cols_sql.append(f"{col} TEXT PRIMARY KEY")
        elif col == "row_no":
            cols_sql.append(f"{col} INTEGER")
        else:
            cols_sql.append(f"{col} TEXT")
    ddl = (
        f"CREATE TABLE IF NOT EXISTS {TABLE} (\n  "
        + ",\n  ".join(cols_sql)
        + ",\n  updated_at TIMESTAMPTZ NOT NULL DEFAULT now()\n);"
    )
    with _connect() as conn, conn.cursor() as cur:
        cur.execute(ddl)
        conn.commit()
    _invalidate()


# ── Read cache (short TTL; invalidated on writes) ───────────────────────────
_CACHE: dict = {"ts": 0.0, "rows": None}
_LOCK = threading.Lock()
_TTL = 1.5  # seconds


def _invalidate() -> None:
    with _LOCK:
        _CACHE["rows"] = None
        _CACHE["ts"] = 0.0


def _row_to_headerdict(dbrow: dict) -> dict:
    out = {}
    for header, col in COLUMNS:
        if col == "row_no":
            out[header] = dbrow.get(col)
        else:
            out[header] = dbrow.get(col)
    return out


def all_rows(use_cache: bool = True) -> list[dict]:
    with _LOCK:
        if use_cache and _CACHE["rows"] is not None and (time.time() - _CACHE["ts"]) < _TTL:
            return _CACHE["rows"]
    cols = ", ".join(col for _, col in COLUMNS)
    with _connect() as conn, conn.cursor() as cur:
        cur.execute(
            f"SELECT {cols} FROM {TABLE} "
            f"ORDER BY row_no NULLS LAST, company_number"
        )
        rows = [_row_to_headerdict(r) for r in cur.fetchall()]
    with _LOCK:
        _CACHE["rows"] = rows
        _CACHE["ts"] = time.time()
    return rows


def find_row(cn: str) -> dict | None:
    target = _cn8(cn)
    for r in all_rows():
        if _cn8(r.get("Company Number")) == target:
            return r
    return None


def company_numbers() -> set[str]:
    return {_cn8(r.get("Company Number")) for r in all_rows() if r.get("Company Number")}


def _next_row_no(cur) -> int:
    cur.execute(f"SELECT COALESCE(MAX(row_no), 0) AS m FROM {TABLE}")
    return int(cur.fetchone()["m"]) + 1


# ── Writes ──────────────────────────────────────────────────────────────────
def add_company(cn: str, name: str = "", date_of_creation: str = "") -> bool:
    """Insert a new company. Returns True if inserted, False if it already exists."""
    cn = _cn8(cn)
    if not cn or cn == "00000000":
        return False
    with _connect() as conn, conn.cursor() as cur:
        cur.execute(f"SELECT 1 FROM {TABLE} WHERE company_number = %s", (cn,))
        if cur.fetchone():
            return False
        rn = _next_row_no(cur)
        cur.execute(
            f"INSERT INTO {TABLE} (company_number, company_name, date_of_creation, row_no) "
            f"VALUES (%s, %s, %s, %s) ON CONFLICT (company_number) DO NOTHING",
            (cn, _norm(name), _norm(date_of_creation), rn),
        )
        conn.commit()
    _invalidate()
    return True


def update_row(cn: str, updates: dict) -> None:
    """Update columns for a company (keys are header labels). Upserts if new."""
    cn = _cn8(cn)
    if not cn:
        return
    set_cols = {}
    for header, value in updates.items():
        col = HEADER_TO_COL.get(header)
        if col is None or col == "company_number":
            continue
        set_cols[col] = _norm(value)
    if not set_cols:
        return
    with _connect() as conn, conn.cursor() as cur:
        cur.execute(f"SELECT 1 FROM {TABLE} WHERE company_number = %s", (cn,))
        exists = cur.fetchone() is not None
        if not exists:
            rn = _next_row_no(cur)
            set_cols.setdefault("row_no", rn)
            cols = ["company_number"] + list(set_cols.keys())
            vals = [cn] + list(set_cols.values())
            placeholders = ", ".join(["%s"] * len(cols))
            cur.execute(
                f"INSERT INTO {TABLE} ({', '.join(cols)}) VALUES ({placeholders})",
                vals,
            )
        else:
            assignments = ", ".join(f"{c} = %s" for c in set_cols)
            cur.execute(
                f"UPDATE {TABLE} SET {assignments}, updated_at = now() "
                f"WHERE company_number = %s",
                list(set_cols.values()) + [cn],
            )
        conn.commit()
    _invalidate()


# Map a nested pipeline result dict (as produced by run_pipeline) to header labels.
def _result_to_headers(r: dict) -> dict:
    d   = r.get("details", {}) or {}
    dn  = r.get("duns", {}) or {}
    ce  = r.get("certificate", {}) or {}
    dm  = r.get("domain", {}) or {}
    em  = r.get("email", {}) or {}
    pc  = r.get("play_console", {}) or {}
    ph  = r.get("phone", {}) or {}
    cert_ok = (ce.get("status") == "downloaded") or bool(ce.get("path"))
    return {
        "Company Number":         r.get("company_number", ""),
        "Company Name":           d.get("company_name", ""),
        "Status":                 d.get("company_status", ""),
        "Type":                   d.get("company_type", ""),
        "Date of Creation":       d.get("date_of_creation", ""),
        "SIC Codes":              d.get("sic_codes", ""),
        "Address":                d.get("address", ""),
        "Directors":              d.get("director_names", ""),
        "Director Nationalities": d.get("director_nationalities", ""),
        "DUNS Number":            dn.get("duns_number", ""),
        "DUNS Status":            dn.get("status", ""),
        "DUNS Email Used":        dn.get("temp_email", ""),
        "D&B Legal Name":         dn.get("dnb_name", ""),
        "D&B Address":            dn.get("dnb_address", ""),
        "Certificate Downloaded": "Yes" if cert_ok else "",
        "Certificate Path":       ce.get("path", "") or ce.get("error", ""),
        "Domain":                 dm.get("domain", ""),
        "Domain Status":          dm.get("status", ""),
        "Domain Cost":            dm.get("charged", ""),
        "Assigned Email":         em.get("email", ""),
        "Account Name":           em.get("account_name", ""),
        "Developer Email":        pc.get("developer_email", ""),
        "Dev Email Forwarding":   (pc.get("developer_email_forwarding", {}) or {}).get("status", ""),
        "Google TXT Status":      (pc.get("google_txt", {}) or {}).get("status", ""),
        "Organization Phone":     ph.get("phone_number", ""),
    }


def _headers_to_result(row: dict) -> dict:
    """Inverse of _result_to_headers: header-keyed row -> nested pipeline result."""
    g = lambda h: row.get(h) or ""
    cert_dl = str(g("Certificate Downloaded")).strip().lower() in ("yes", "true", "1")
    return {
        "company_number": str(g("Company Number")),
        "details": {
            "company_name":           g("Company Name"),
            "company_status":         g("Status"),
            "company_type":           g("Type"),
            "date_of_creation":       g("Date of Creation"),
            "sic_codes":              g("SIC Codes"),
            "address":                g("Address"),
            "director_names":         g("Directors"),
            "director_nationalities": g("Director Nationalities"),
        },
        "duns": {
            "duns_number": g("DUNS Number"),
            "status":      g("DUNS Status"),
            "temp_email":  g("DUNS Email Used"),
            "dnb_name":    g("D&B Legal Name"),
            "dnb_address": g("D&B Address"),
        },
        "certificate": {
            "path":   g("Certificate Path") if cert_dl else "",
            "status": "downloaded" if cert_dl else "",
            "error":  g("Certificate Path") if not cert_dl else "",
        },
        "domain":  {"domain": g("Domain"), "status": g("Domain Status"), "charged": g("Domain Cost")},
        "email":   {"email": g("Assigned Email"), "account_name": g("Account Name")},
        "play_console": {
            "developer_email": g("Developer Email"),
            "developer_email_forwarding": {"status": g("Dev Email Forwarding")},
            "google_txt": {"status": g("Google TXT Status")},
        },
        "phone": {"phone_number": g("Organization Phone")},
    }


def all_results() -> list[dict]:
    """All companies as nested pipeline-result dicts (used by run_pipeline)."""
    return [_headers_to_result(r) for r in all_rows(use_cache=False) if r.get("Company Number")]


def save_results(results: list[dict]) -> int:
    """Upsert a list of nested pipeline result dicts. Returns count processed."""
    n = 0
    for r in results:
        cn = _cn8(r.get("company_number", ""))
        if not cn or cn == "00000000":
            continue
        headers = {k: v for k, v in _result_to_headers(r).items() if v not in (None, "")}
        headers.pop("Company Number", None)
        update_row(cn, headers)
        n += 1
    return n


# ── One-time migration from the old Excel file ──────────────────────────────
def import_from_excel(path: str) -> int:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    hdr_row = next(it, None)
    if not hdr_row:
        wb.close()
        return 0
    hdr = [str(v or "").strip() for v in hdr_row]
    imported = 0
    with _connect() as conn, conn.cursor() as cur:
        for row in it:
            if not row or not row[1 if len(row) > 1 else 0]:
                continue
            rec = {hdr[i]: (row[i] if i < len(row) else None) for i in range(len(hdr))}
            cn = _cn8(rec.get("Company Number"))
            if not cn or cn == "00000000":
                continue
            set_cols = {"company_number": cn}
            for header, col in COLUMNS:
                if col == "company_number":
                    continue
                if header in rec:
                    set_cols[col] = _norm(rec.get(header))
            cols = list(set_cols.keys())
            vals = list(set_cols.values())
            placeholders = ", ".join(["%s"] * len(cols))
            updates = ", ".join(f"{c} = EXCLUDED.{c}" for c in cols if c != "company_number")
            cur.execute(
                f"INSERT INTO {TABLE} ({', '.join(cols)}) VALUES ({placeholders}) "
                f"ON CONFLICT (company_number) DO UPDATE SET {updates}",
                vals,
            )
            imported += 1
        conn.commit()
    wb.close()
    _invalidate()
    return imported


if __name__ == "__main__":
    import sys
    ensure_schema()
    print(f"Schema ready. Table '{TABLE}' has {len(COLUMNS)} mapped columns.")
    if len(sys.argv) > 1 and sys.argv[1] == "--import":
        path = sys.argv[2] if len(sys.argv) > 2 else "pipeline_output/companies_pipeline.xlsx"
        n = import_from_excel(path)
        print(f"Imported {n} rows from {path}.")
    print(f"Total rows in DB: {len(all_rows(use_cache=False))}")
