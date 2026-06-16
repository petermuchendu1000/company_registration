"""Load the company catalog from the processed pipeline Excel."""

from __future__ import annotations

import os
import re
from dataclasses import dataclass
from typing import Iterable

import openpyxl

EXCEL_PATH_DEFAULT = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "pipeline_output",
    "companies_pipeline.xlsx",
)

# SIC code -> archetype. Only "shift" exists today; the rest are placeholders
# for future archetype modules and currently fall back to "shift".
SIC_ARCHETYPE = {
    "78200": "shift",    # Temp employment agency
    "88100": "shift",    # Social work without accommodation for the elderly/disabled
    "87100": "shift",    # Residential nursing care
    "53202": "shift",    # Courier (will become "mileage" when that archetype lands)
    "49410": "shift",    # Freight road transport
    "98000": "shift",    # Residents property mgmt (will become "twa" when that lands)
}
DEFAULT_ARCHETYPE = "shift"


DISPLAY_NAME_LIMIT = 30
LEGAL_SUFFIX_RE = re.compile(r"\s+(LTD|LIMITED|LLP|PLC)\.?$", re.IGNORECASE)
DISPLAY_ABBREVIATIONS = {
    "Association": "Assoc",
    "Company": "Co",
    "Development": "Dev",
    "Developments": "Dev",
    "International": "Intl",
    "Management": "Mgmt",
    "Properties": "Props",
    "Property": "Prop",
    "Services": "Svcs",
}


def _title_case_name(value: str) -> str:
    name = value.strip()
    name = LEGAL_SUFFIX_RE.sub("", name)
    name = re.sub(r"\s+", " ", name)
    return name.title()


def _fit_display_name(value: str) -> str:
    """Return a Play-safe title without chopping through a word."""
    name = _title_case_name(value)
    if len(name) <= DISPLAY_NAME_LIMIT:
        return name

    abbreviated = name
    for source, replacement in DISPLAY_ABBREVIATIONS.items():
        abbreviated = re.sub(rf"\b{source}\b", replacement, abbreviated)
        if len(abbreviated) <= DISPLAY_NAME_LIMIT:
            return abbreviated

    words = abbreviated.split()
    while len(" ".join(words)) > DISPLAY_NAME_LIMIT and len(words) > 1:
        words.pop()
    shortened = " ".join(words).strip(" ,.-")
    if shortened:
        return shortened
    return abbreviated[:DISPLAY_NAME_LIMIT].rstrip(" ,.-")


@dataclass(frozen=True)
class Company:
    company_number: str
    company_name: str
    sic_codes: str
    domain: str
    support_email: str
    archetype: str
    developer_display_name: str = ""  # Human name from Gmail account (e.g. "Faith Kabebee")
    organization_phone: str = ""  # Org phone from phone pool (e.g. "+254712345678")

    @property
    def flavor(self) -> str:
        """Gradle flavor name. Keep it lowercase for predictable task names."""
        token = re.sub(r"[^a-z0-9]", "", self.company_number.lower())
        return f"c{token}"

    @property
    def application_id(self) -> str:
        token = re.sub(r"[^a-z0-9]", "", self.company_number.lower())
        return f"uk.c{token}.shift"

    @property
    def display_name(self) -> str:
        """Launcher label with legal suffix stripped and Google Play's 30-char cap."""
        return _fit_display_name(self.company_name)


def load_companies(
    excel_path: str = EXCEL_PATH_DEFAULT,
    limit: int | None = None,
) -> list[Company]:
    """Read companies from the pipeline Excel. Skips rows without a company number."""
    wb = openpyxl.load_workbook(excel_path, data_only=True, read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows)

    def col(name: str) -> int:
        for i, h in enumerate(header):
            if h and h.strip().lower() == name.lower():
                return i
        raise KeyError(f"column not found: {name}")

    def col_opt(name: str) -> int | None:
        """Like col() but returns None if the column is absent."""
        for i, h in enumerate(header):
            if h and h.strip().lower() == name.lower():
                return i
        return None

    idx_num = col("Company Number")
    idx_name = col("Company Name")
    idx_sic = col("SIC Codes")
    idx_domain = col("Domain")
    idx_email = col("Assigned Email")
    idx_dev_name  = col_opt("Account Name")      # Gmail display name
    idx_org_phone = col_opt("Organization Phone") # Phone pool

    out: list[Company] = []
    for row in rows:
        if not row or not row[idx_num]:
            continue
        cn = str(row[idx_num]).strip()
        # Pad to 8 digits (Companies House format)
        if cn.isdigit() and len(cn) < 8:
            cn = cn.zfill(8)

        sic_raw = (row[idx_sic] or "")
        primary_sic = sic_raw.split(",")[0].strip() if sic_raw else ""
        archetype = SIC_ARCHETYPE.get(primary_sic, DEFAULT_ARCHETYPE)

        domain = (row[idx_domain] or "").strip() if row[idx_domain] else ""
        email = (row[idx_email] or "").strip() if row[idx_email] else ""
        support = f"dev@{domain}" if domain else (email or "support@example.uk")

        out.append(Company(
            company_number=cn,
            company_name=(row[idx_name] or "").strip(),
            sic_codes=sic_raw,
            domain=domain,
            support_email=support,
            archetype=archetype,
            developer_display_name=(
                str(row[idx_dev_name]).strip()
                if idx_dev_name is not None and idx_dev_name < len(row) and row[idx_dev_name]
                else ""
            ),
            organization_phone=(
                str(row[idx_org_phone]).strip()
                if idx_org_phone is not None and idx_org_phone < len(row) and row[idx_org_phone]
                else ""
            ),
        ))
        if limit is not None and len(out) >= limit:
            break

    wb.close()
    return out


def summarize(companies: Iterable[Company]) -> str:
    lines = [f"{'#':>3}  {'Flavor':<12}  {'SIC':<6}  {'Arch':<6}  Name"]
    for i, c in enumerate(companies, 1):
        primary_sic = c.sic_codes.split(",")[0].strip() if c.sic_codes else "-"
        lines.append(f"{i:>3}  {c.flavor:<12}  {primary_sic:<6}  {c.archetype:<6}  {c.display_name}")
    return "\n".join(lines)


if __name__ == "__main__":
    import sys
    n = int(sys.argv[1]) if len(sys.argv) > 1 else 3
    cs = load_companies(limit=n)
    print(summarize(cs))
