"""Load the company catalog from the processed pipeline Excel."""
from __future__ import annotations
import os
import re
from dataclasses import dataclass
from typing import Iterable
import openpyxl

EXCEL_PATH_DEFAULT = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "pipeline_output",
    "companies_pipeline.xlsx",
)

# SIC code -> archetype. Only "shift" exists today; the rest are placeholders
# for future archetype modules and currently fall back to "shift".
SIC_ARCHETYPE = {
    "78200": "shift",  # Temp employment agency
    "88100": "shift",  # Social work without accommodation for the elderly/disabled
    "87100": "shift",  # Residential nursing care
    "53202": "shift",  # Courier (will become "mileage" when that archetype lands)
    "49410": "shift",  # Freight road transport
    "98000": "shift",  # Residents property mgmt (will become "twa" when that lands)
}
DEFAULT_ARCHETYPE = "shift"


@dataclass(frozen=True)
class Company:
    company_number: str
    company_name: str
    sic_codes: str
    domain: str
    support_email: str
    archetype: str
    short_name: str = ""
    address: str = ""

    @property
    def flavor(self) -> str:
        return f"c{self.company_number}"

    @property
    def slug(self) -> str:
        """Unique lowercase slug from company name — used as applicationId segment."""
        name = re.sub(r"\s+(LTD|LIMITED|LLP|PLC|INC)\.?$", "", self.company_name, flags=re.IGNORECASE)
        for stop in ("management", "services", "solutions", "group", "holdings",
                     "enterprises", "consulting", "associates", "the", "and", "of"):
            name = re.sub(rf"\b{stop}\b", "", name, flags=re.IGNORECASE)
        slug = re.sub(r"[^a-zA-Z0-9]", "", name).lower()
        if not slug or not slug[0].isalpha():
            slug = "co" + slug
        return slug[:20] or f"co{self.company_number}"

    @property
    def application_id(self) -> str:
        return f"uk.{self.slug}.app"

    @property
    def display_name(self) -> str:
        """Launcher label — prefers curated short_name; falls back to
        title-cased company name with LTD/LIMITED stripped, 30-char cap
        that breaks on a word boundary to avoid mid-word cut-offs."""
        if self.short_name:
            return self.short_name[:30].rstrip()

        name = self.company_name.strip()
        name = re.sub(r"\s+(LTD|LIMITED|LLP|PLC)\.?$", "", name, flags=re.IGNORECASE)
        name = name.title()
        if len(name) <= 30:
            return name
        # Truncate at last space within the 30-char window
        truncated = name[:30]
        last_space = truncated.rfind(" ")
        return (truncated[:last_space] if last_space > 0 else truncated).rstrip()


def load_companies(
    excel_path: str = EXCEL_PATH_DEFAULT,
    limit: int | None = None,
    company_numbers: list[str] | None = None,
) -> list[Company]:
    """Read companies from the pipeline Excel. Skips rows without a company number."""
    wb = openpyxl.load_workbook(excel_path, data_only=True, read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows)

    def col(name: str) -> int:
        for i, h in enumerate(header):
            if h and h.strip().lower() == name.lower():
                return i
        raise KeyError(f"column not found: {name}")

    def col_optional(name: str) -> int | None:
        """Return column index or None if the column doesn't exist yet."""
        for i, h in enumerate(header):
            if h and h.strip().lower() == name.lower():
                return i
        return None

    idx_num    = col("Company Number")
    idx_name   = col("Company Name")
    idx_sic    = col("SIC Codes")
    idx_domain = col("Domain")
    idx_email  = col("Assigned Email")
    idx_short  = col_optional("Short Name")
    idx_addr   = col_optional("Address")

    # Normalise filter set once
    filter_set = {n.zfill(8) if n.isdigit() and len(n) < 8 else n for n in (company_numbers or [])}

    out: list[Company] = []
    for row in rows:
        if not row or not row[idx_num]:
            continue

        cn = str(row[idx_num]).strip()
        # Pad to 8 digits (Companies House format)
        if cn.isdigit() and len(cn) < 8:
            cn = cn.zfill(8)

        if filter_set and cn not in filter_set:
            continue

        sic_raw = (row[idx_sic] or "")
        primary_sic = sic_raw.split(",")[0].strip() if sic_raw else ""
        archetype = SIC_ARCHETYPE.get(primary_sic, DEFAULT_ARCHETYPE)

        domain = (row[idx_domain] or "").strip() if row[idx_domain] else ""
        email  = (row[idx_email] or "").strip() if row[idx_email] else ""
        support = f"dev@{domain}" if domain else (email or "support@example.uk")

        short = ""
        if idx_short is not None:
            short = (row[idx_short] or "").strip() if row[idx_short] else ""

        addr = ""
        if idx_addr is not None:
            addr = (row[idx_addr] or "").strip() if row[idx_addr] else ""

        out.append(Company(
            company_number=cn,
            company_name=(row[idx_name] or "").strip(),
            sic_codes=sic_raw,
            domain=domain,
            support_email=support,
            archetype=archetype,
            short_name=short,
            address=addr,
        ))

        if limit is not None and len(out) >= limit:
            break

    wb.close()
    return out


def summarize(companies: Iterable[Company]) -> str:
    lines = [f"{'#':>3} {'Flavor':<12} {'SIC':<6} {'Arch':<6} Name"]
    for i, c in enumerate(companies, 1):
        primary_sic = c.sic_codes.split(",")[0].strip() if c.sic_codes else "-"
        lines.append(f"{i:>3} {c.flavor:<12} {primary_sic:<6} {c.archetype:<6} {c.display_name}")
    return "\n".join(lines)


if __name__ == "__main__":
    import sys
    n = int(sys.argv[1]) if len(sys.argv) > 1 else 3
    cs = load_companies(limit=n)
    print(summarize(cs))
